"""OSL parser: a Word document read by heading and table.

Implements :class:`greenlight_ai.parsers.base.OslParser` (ADR-006). The OSL is the source of
truth (ADR-002), so this parser preserves text verbatim: it never normalises wording,
expands abbreviations, or drops a sentence it does not understand. Interpretation is
stage 2's job, and it is done by the model, not here.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Final, Iterator

from greenlight_ai.parsers.base import OslDocument, OslSection, OslTable, ParseError

__all__ = ["DocxOslParser"]

_LOG: Final = logging.getLogger(__name__)

#: Matches a leading section number such as "3", "3.1", or "3.1.2", with optional
#: trailing punctuation, so "3.1 Geography" and "3.1. Geography" both split correctly.
_NUMBER_RE: Final = re.compile(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+(.*)$")

#: Heading style names python-docx reports, lowercased. "Title" is treated as the
#: document title rather than a section.
_HEADING_PREFIX: Final = "heading"
_TITLE_STYLE: Final = "title"

#: Text placed in a section that has no heading of its own (content before the first
#: heading, which real specs do carry: scope notes, revision tables).
_PREAMBLE_NUMBER: Final = "0"
_PREAMBLE_HEADING: Final = "Preamble"


class DocxOslParser:
    """Parses a ``.docx`` requirement spec into sections and tables."""

    def parse(self, path: Path) -> OslDocument:
        """Parse the OSL at ``path``.

        Args:
            path: The ``.docx`` file to read.

        Returns:
            The parsed document, sections in document order, each carrying the paragraphs
            and tables that follow its heading.

        Raises:
            ParseError: When the file is missing, is not a ``.docx``, or python-docx
                cannot open it.
        """
        try:
            import docx  # noqa: PLC0415 - imported lazily so the package imports without it
        except ImportError as exc:  # pragma: no cover - dependency is declared
            raise ParseError(path, "python-docx is not installed") from exc

        if not path.exists():
            raise ParseError(path, "file does not exist")
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # noqa: BLE001 - python-docx raises bare exceptions
            raise ParseError(path, f"could not be opened as a Word document ({exc})") from exc

        title = ""
        sections: list[OslSection] = []
        number, heading, level = _PREAMBLE_NUMBER, _PREAMBLE_HEADING, 1
        paragraphs: list[str] = []
        tables: list[OslTable] = []

        for item in _iter_body(document):
            if isinstance(item, _Heading):
                if paragraphs or tables or sections:
                    sections.append(
                        OslSection(
                            number=number,
                            heading=heading,
                            level=level,
                            paragraphs=tuple(paragraphs),
                            tables=tuple(tables),
                        )
                    )
                number, heading, level = item.number, item.text, item.level
                paragraphs, tables = [], []
            elif isinstance(item, _Title):
                title = item.text
            elif isinstance(item, _Paragraph):
                paragraphs.append(item.text)
            else:
                tables.append(OslTable(index=len(tables) + 1, header=item.header, rows=item.rows))

        sections.append(
            OslSection(
                number=number,
                heading=heading,
                level=level,
                paragraphs=tuple(paragraphs),
                tables=tuple(tables),
            )
        )
        sections = [s for s in sections if s.paragraphs or s.tables or s.number != _PREAMBLE_NUMBER]

        _LOG.info(
            "parsed OSL %s: %d sections, %d tables",
            path.name,
            len(sections),
            sum(len(s.tables) for s in sections),
        )
        return OslDocument(path=path, title=title or path.stem, sections=tuple(sections))


# --------------------------------------------------------------------------------------
# Body iteration. python-docx exposes paragraphs and tables as separate collections, so
# document order is recovered from the underlying XML.
# --------------------------------------------------------------------------------------


class _Heading:
    """A heading encountered in the body."""

    __slots__ = ("number", "text", "level")

    def __init__(self, number: str, text: str, level: int) -> None:
        self.number = number
        self.text = text
        self.level = level


class _Title:
    """The document title."""

    __slots__ = ("text",)

    def __init__(self, text: str) -> None:
        self.text = text


class _Paragraph:
    """A body paragraph."""

    __slots__ = ("text",)

    def __init__(self, text: str) -> None:
        self.text = text


class _Table:
    """A table, already flattened to text."""

    __slots__ = ("header", "rows")

    def __init__(self, header: tuple[str, ...], rows: tuple[tuple[str, ...], ...]) -> None:
        self.header = header
        self.rows = rows


def _iter_body(document: object) -> Iterator[_Heading | _Title | _Paragraph | _Table]:
    """Walk a document's body in true document order.

    Args:
        document: The python-docx ``Document``.

    Yields:
        Headings, the title, paragraphs, and tables, in the order they appear.
    """
    body = document.element.body  # type: ignore[attr-defined]
    paragraphs = {p._element: p for p in document.paragraphs}  # type: ignore[attr-defined]
    tables = {t._element: t for t in document.tables}  # type: ignore[attr-defined]

    for child in body.iterchildren():
        if child in paragraphs:
            paragraph = paragraphs[child]
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").strip().lower()
            if style == _TITLE_STYLE:
                yield _Title(text)
            elif style.startswith(_HEADING_PREFIX):
                level = _heading_level(style)
                number, heading = _split_number(text)
                yield _Heading(number, heading, level)
            else:
                yield _Paragraph(text)
        elif child in tables:
            header, rows = _flatten(tables[child])
            if header or rows:
                yield _Table(header, rows)


def _heading_level(style: str) -> int:
    """Extract the depth from a heading style name.

    Args:
        style: The lowercased style name, e.g. ``"heading 2"``.

    Returns:
        The depth, defaulting to 1 when the name carries no digit.
    """
    match = re.search(r"(\d+)", style)
    return int(match.group(1)) if match else 1


def _split_number(text: str) -> tuple[str, str]:
    """Split a heading into its number and its text.

    Args:
        text: The heading as written, e.g. ``"3.1 Geography"``.

    Returns:
        A ``(number, heading)`` pair. When the heading carries no number the number is
        the empty string, which keeps ``source_ref`` honest rather than inventing one.
    """
    match = _NUMBER_RE.match(text)
    if match:
        return match.group(1), match.group(2).strip()
    return "", text


def _flatten(table: object) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    """Flatten a python-docx table to header and rows.

    Merged cells repeat their text in python-docx; duplicates are left as-is because the
    model reads the table and the repetition is information, not noise.

    Args:
        table: The python-docx ``Table``.

    Returns:
        A ``(header, rows)`` pair. Ragged rows are padded to the header width.
    """
    grid: list[tuple[str, ...]] = []
    for row in table.rows:  # type: ignore[attr-defined]
        grid.append(tuple(cell.text.strip() for cell in row.cells))
    if not grid:
        return (), ()
    header = grid[0]
    width = len(header)
    rows = tuple(row + ("",) * (width - len(row)) if len(row) < width else row for row in grid[1:])
    return header, rows
