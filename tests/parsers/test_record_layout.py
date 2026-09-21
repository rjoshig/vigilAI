"""Reading the delivered file's record schema (Phase 6.22b).

The rules under test are the three the module docstring states: the layout is optional,
its own headers go up the ladder, and it carries names and shapes rather than rows. The
second is the one with teeth — every real layout workbook heads its columns differently,
and a parser that insisted on one spelling would report a delivery as declaring no
fields at all, which is the loudest possible wrong answer.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from greenlight_ai.parsers.base import ParseError
from greenlight_ai.parsers.record_layout import (
    RecordLayoutDocument,
    RecordLayoutField,
    parse_record_layout,
)

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


def _write(path: Path, header: list[str], rows: list[list[object]], title: str = "Layout") -> Path:
    """Write a one-sheet layout workbook."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


class TestReadingOne:
    """What a layout workbook parses into."""

    def test_the_three_columns_are_read_in_record_order(self, tmp_path: Path) -> None:
        path = _write(
            tmp_path / "layout.xlsx",
            ["Field name", "Data type", "Size"],
            [["SCORE_V3", "DECIMAL", 4], ["ST", "CHAR", 2]],
        )
        document = parse_record_layout(path)
        assert document.names == ("SCORE_V3", "ST")
        assert [f.ordinal for f in document.fields] == [1, 2]
        assert document.fields[0].shape == "DECIMAL(4)"

    @pytest.mark.parametrize(
        "header",
        [
            ["Column Name", "Type", "Length"],
            ["field_name", "data_type", "size"],
            ["ATTRIBUTE NAME", "Field type", "Max length"],
        ],
    )
    def test_a_layout_headed_differently_is_the_same_document(
        self, tmp_path: Path, header: list[str]
    ) -> None:
        """The ladder, not a fixed spelling. This is the whole reason 6.21a exists."""
        path = _write(tmp_path / "layout.xlsx", header, [["SCORE_V3", "DECIMAL", 4]])
        document = parse_record_layout(path)
        assert document.names == ("SCORE_V3",)
        assert document.fields[0].data_type == "DECIMAL"
        assert document.fields[0].size == "4"

    def test_a_size_excel_stored_as_a_float_reads_as_the_integer_the_layout_wrote(
        self, tmp_path: Path
    ) -> None:
        """A layout saying 10 must not be read back as 10.0 and quoted that way."""
        path = _write(
            tmp_path / "layout.xlsx", ["Field name", "Data type", "Size"], [["ST", "CHAR", 10.0]]
        )
        assert parse_record_layout(path).fields[0].size == "10"

    def test_a_cover_sheet_before_the_layout_is_skipped(self, tmp_path: Path) -> None:
        """Real layout workbooks open with a title page. That is not a parse failure."""
        import openpyxl

        path = tmp_path / "layout.xlsx"
        workbook = openpyxl.Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover.append(["Prepared for", "A Customer"])
        cover.append(["Version", "3"])
        sheet = workbook.create_sheet("Fields")
        sheet.append(["Field name", "Data type", "Size"])
        sheet.append(["AT01", "CHAR", 6])
        workbook.save(path)

        assert parse_record_layout(path).names == ("AT01",)

    def test_the_type_and_size_columns_are_optional(self, tmp_path: Path) -> None:
        """A layout that lists names and nothing else is still a layout."""
        path = _write(tmp_path / "layout.xlsx", ["Field name"], [["AT01"], ["ST"]])
        document = parse_record_layout(path)
        assert document.names == ("AT01", "ST")
        assert document.fields[0].shape == ""

    def test_a_workbook_with_no_field_name_column_is_a_parse_error(self, tmp_path: Path) -> None:
        """Not an empty layout.

        Reading it as one would silently answer "this delivery declares no fields",
        which is a false assertion rather than a missing answer — the exact conflation
        6.22a was written to undo.
        """
        path = _write(tmp_path / "layout.xlsx", ["Sheet", "Rows"], [["Summary", 12]])
        with pytest.raises(ParseError, match="field-name column"):
            parse_record_layout(path)

    def test_a_missing_file_is_a_parse_error(self, tmp_path: Path) -> None:
        with pytest.raises(ParseError, match="does not exist"):
            parse_record_layout(tmp_path / "nope.xlsx")

    def test_a_blank_name_does_not_become_a_field(self, tmp_path: Path) -> None:
        """A spacer row in the middle of a layout is a spacer, not a field called ''."""
        path = _write(
            tmp_path / "layout.xlsx",
            ["Field name", "Data type", "Size"],
            [["AT01", "CHAR", 6], [None, None, None], ["ST", "CHAR", 2]],
        )
        document = parse_record_layout(path)
        assert document.names == ("AT01", "ST")
        assert [f.ordinal for f in document.fields] == [1, 2]


class TestLookingSomethingUp:
    """Resolving a name against the layout goes up the same ladder as everything else."""

    def _document(self) -> RecordLayoutDocument:
        return RecordLayoutDocument(
            fields=(
                RecordLayoutField("SCORE_V3", "DECIMAL", "4", 1),
                RecordLayoutField("opt_out", "CHAR", "1", 2),
            )
        )

    def test_an_exact_name_resolves(self) -> None:
        assert self._document().carries("SCORE_V3")

    def test_separators_are_noise(self) -> None:
        """Rung 2. ``OPT-OUT`` and ``opt_out`` are one name."""
        assert self._document().carries("OPT-OUT")

    def test_a_name_nothing_resembles_does_not_resolve(self) -> None:
        assert not self._document().carries("HOUSEHOLD_INCOME")
        assert self._document().field("HOUSEHOLD_INCOME") is None


class TestWhereItCameFrom:
    """A borrowed layout says so, in the words a finding carries."""

    def test_a_layout_this_run_uploaded_is_not_borrowed(self) -> None:
        document = RecordLayoutDocument(fields=(RecordLayoutField("ST"),))
        assert not document.borrowed
        assert document.provenance == ""

    def test_a_borrowed_layout_names_the_run_and_the_date(self) -> None:
        document = RecordLayoutDocument(
            fields=(RecordLayoutField("ST"),), source_run_id=41, source_date="2026-08-14"
        )
        assert document.borrowed
        assert document.provenance == "from run 41 finalized 2026-08-14"

    def test_a_borrowed_layout_with_no_date_still_names_the_run(self) -> None:
        document = RecordLayoutDocument(fields=(RecordLayoutField("ST"),), source_run_id=41)
        assert document.provenance == "from run 41"


class TestTheStoredShape:
    """A snapshot round-trips, and a strange one degrades rather than raising."""

    def test_as_rows_and_from_rows_round_trip(self) -> None:
        document = RecordLayoutDocument(
            fields=(
                RecordLayoutField("ST", "CHAR", "2", 1),
                RecordLayoutField("AGE", "INT", "3", 2),
            )
        )
        rebuilt = RecordLayoutDocument.from_rows(document.as_rows())
        assert rebuilt.fields == document.fields

    @pytest.mark.parametrize("stored", [None, "", 12, [], [1, 2], [{"size": "4"}], [{"name": " "}]])
    def test_a_snapshot_that_has_gone_strange_gives_an_empty_layout(self, stored: object) -> None:
        """A run must not fail because a stored layout is not what it was.

        An empty layout is the ordinary state, so degrading to it changes nothing about
        how the run is checked.
        """
        assert RecordLayoutDocument.from_rows(stored).fields == ()

    def test_a_row_with_no_ordinal_gets_its_position(self) -> None:
        rebuilt = RecordLayoutDocument.from_rows([{"name": "ST"}, {"name": "AGE"}])
        assert [f.ordinal for f in rebuilt.fields] == [1, 2]


class TestTheFixture:
    """The generated case, read the way the pipeline reads it."""

    def test_the_shipped_case_parses_through_its_drifted_headers(self) -> None:
        path = FIXTURES / "cases/record_layout_supplied/record_layout.xlsx"
        document = parse_record_layout(path)
        assert document.carries("SCORE_V3")
        # Two fields nobody asked for, which 6.22c turns into a low-severity note.
        assert document.carries("INTERNAL_SEQ")
        assert len(document.fields) == 10
